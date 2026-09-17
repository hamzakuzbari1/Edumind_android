"""Export parent historical reports — CSV, Excel, PDF."""

from __future__ import annotations

import csv
import io
from typing import Any


def _metric_row(label: str, metric: dict | None) -> list[str]:
    if not metric:
        return [label, "—", "—", "—"]
    change = metric.get("change_percent")
    change_str = f"{change:+.1f}%" if change is not None else "—"
    unit = metric.get("unit") or ""
    cur = metric.get("current_value")
    prev = metric.get("previous_value")
    return [
        label,
        f"{cur} {unit}".strip(),
        f"{prev} {unit}".strip(),
        change_str,
    ]


def export_report_csv(report: dict[str, Any]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["تقرير ولي الأمر — EduSpark"])
    writer.writerow(["الطالب", report.get("student_name", "")])
    writer.writerow(["الصف", report.get("grade_label", "")])
    writer.writerow(["الفترة", report.get("period_label", "")])
    writer.writerow(["من", report.get("start_date", ""), "إلى", report.get("end_date", "")])
    writer.writerow([])

    cmp_ = report.get("comparison") or {}
    writer.writerow(["مقارنة الفترات", cmp_.get("period_label", ""), "مقابل", cmp_.get("previous_period_label", "")])
    writer.writerow(["المؤشر", "الحالي", "السابق", "الفرق"])
    for key, ar_label in [
        ("study_time", "وقت الدراسة"),
        ("grades", "متوسط الدرجات"),
        ("lesson_completion", "إكمال الدروس"),
        ("planner_adherence", "الالتزام بالخطة"),
    ]:
        writer.writerow(_metric_row(ar_label, cmp_.get(key)))

    writer.writerow([])
    writer.writerow(["إكمال الدروس — أسبوعياً"])
    writer.writerow(["الفترة", "العدد"])
    for pt in (report.get("lesson_history") or {}).get("weekly") or []:
        writer.writerow([pt.get("label"), pt.get("value")])

    writer.writerow([])
    writer.writerow(["إكمال الدروس — شهرياً"])
    writer.writerow(["الشهر", "العدد"])
    for pt in (report.get("lesson_history") or {}).get("monthly") or []:
        writer.writerow([pt.get("label"), pt.get("value")])

    writer.writerow([])
    writer.writerow(["الحضور — جلسات الدخول"])
    writer.writerow(["التاريخ", "دخول", "خروج", "دقائق نشطة"])
    for s in (report.get("attendance_history") or {}).get("login_sessions") or []:
        writer.writerow(
            [
                s.get("date"),
                s.get("login_at"),
                s.get("logout_at") or "—",
                s.get("active_minutes"),
            ]
        )

    writer.writerow([])
    writer.writerow(["المخطط — الالتزام أسبوعياً"])
    writer.writerow(["الأسبوع", "الالتزام %"])
    for pt in (report.get("planner_history") or {}).get("adherence_trend") or []:
        writer.writerow([pt.get("label"), pt.get("value")])

    writer.writerow([])
    writer.writerow(["المخطط — مهام فائتة/متأخرة"])
    writer.writerow(["الأسبوع", "العدد"])
    for pt in (report.get("planner_history") or {}).get("missed_tasks_trend") or []:
        writer.writerow([pt.get("label"), pt.get("value")])

    return buf.getvalue().encode("utf-8-sig")


def export_report_xlsx(report: dict[str, Any]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError("openpyxl غير مثبت") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "ملخص"
    bold = Font(bold=True)

    ws.append(["تقرير ولي الأمر — EduSpark"])
    ws.append(["الطالب", report.get("student_name", "")])
    ws.append(["الفترة", report.get("period_label", "")])
    ws.append([])

    cmp_ = report.get("comparison") or {}
    ws.append(["المؤشر", "الحالي", "السابق", "الفرق"])
    for key, ar_label in [
        ("study_time", "وقت الدراسة"),
        ("grades", "متوسط الدرجات"),
        ("lesson_completion", "إكمال الدروس"),
        ("planner_adherence", "الالتزام بالخطة"),
    ]:
        ws.append(_metric_row(ar_label, cmp_.get(key)))

    ws2 = wb.create_sheet("الدروس")
    ws2.append(["أسبوع", "دروس مكتملة"])
    for pt in (report.get("lesson_history") or {}).get("weekly") or []:
        ws2.append([pt.get("label"), pt.get("value")])

    ws3 = wb.create_sheet("الحضور")
    ws3.append(["التاريخ", "دخول", "خروج", "دقائق"])
    for s in (report.get("attendance_history") or {}).get("login_sessions") or []:
        ws3.append([s.get("date"), s.get("login_at"), s.get("logout_at"), s.get("active_minutes")])

    ws4 = wb.create_sheet("المخطط")
    ws4.append(["أسبوع", "الالتزام %", "مهام متأخرة"])
    adherence = (report.get("planner_history") or {}).get("adherence_trend") or []
    missed = (report.get("planner_history") or {}).get("missed_tasks_trend") or []
    for i, pt in enumerate(adherence):
        miss_val = missed[i]["value"] if i < len(missed) else 0
        ws4.append([pt.get("label"), pt.get("value"), miss_val])

    for sheet in wb.worksheets:
        sheet["A1"].font = bold

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_report_pdf(report: dict[str, Any]) -> bytes:
    from app.services.parent_report_pdf_builder import build_dashboard_pdf

    return build_dashboard_pdf(report)


def export_filename(report: dict[str, Any], ext: str) -> str:
    """ASCII-safe filename for Content-Disposition (Starlette encodes headers as latin-1)."""
    period = report.get("period") or "report"
    raw_name = (report.get("student_name") or "student").split()[0]
    slug = "".join(
        ch if ch.isascii() and (ch.isalnum() or ch in "-_") else "_"
        for ch in raw_name
    ).strip("_") or "student"
    return f"eduspark-report-{slug}-{period}.{ext}"


def content_disposition_attachment(filename: str) -> dict[str, str]:
    """RFC 5987 — ASCII fallback plus UTF-8 filename* for download clients."""
    from urllib.parse import quote

    ascii_name = filename.encode("ascii", "ignore").decode() or "eduspark-report.csv"
    encoded = quote(filename, safe="")
    return {
        "Content-Disposition": f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{encoded}'
    }
